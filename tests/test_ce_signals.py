import unittest
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

from ocs.ce_signals import booking_plan, fill_ce_signals, item_quantity, load_rules, scaled_quantity
from tp_app.ce_signals_settings import settings_frames, settings_overrides


ROOT = Path(__file__).resolve().parents[1]


def reference_schedule(shift=0):
    from datetime import timedelta
    entries = {
        date(2026, 9, 28): "ST Preparation",
        date(2026, 10, 6): "LMG F2F Clarification",
        date(2026, 10, 9): "Prep for XAW - Stores",
        date(2026, 10, 13): "UO Drills/ SDP",
        date(2026, 10, 14): "PTCO IPSM/ PTCO Truncheon",
        date(2026, 10, 19): "LMG Qualification Shoot",
        date(2026, 10, 25): "EX Relentless Warrior",
        date(2026, 11, 6): "GPMG L/F",
        date(2026, 11, 22): "EX LANCER DAY 7",
        date(2026, 12, 17): "Return of NGLS/Dekitting",
        date(2026, 12, 18): "Return of NGLS/Dekitting",
        date(2026, 12, 20): "Home Sweet Home",
    }
    return pd.DataFrame({(d + timedelta(days=shift)).strftime("%d-%b-%y"): [text]
                         for d, text in entries.items()})


class CESignalsTests(unittest.TestCase):
    def test_legacy_rejects_overwriting_template(self):
        from ocs.legacy_importer import Extractor
        extractor = Extractor.__new__(Extractor)
        extractor.siao_template_path = str(ROOT / "ocs/template_siao.xlsx")
        with self.assertRaisesRegex(ValueError, "separate copy"):
            extractor.draft_siao(140, output_path=extractor.siao_template_path)

    def test_settings_roundtrip_and_generation(self):
        defaults = load_rules()
        items, windows = settings_frames(defaults)
        self.assertEqual(settings_overrides(defaults, 140, items, windows), {})
        items.loc[13, "Baseline quantity"] = 30
        items.loc[15, "Quantity rule"] = "Fixed"
        windows.loc["diesel", "End offset (days)"] = 5
        overrides = settings_overrides(defaults, 140, items, windows)
        rules = load_rules(overrides=overrides)
        _, bookings, _ = booking_plan(reference_schedule(), 280, rules)
        by_name = {b["name"]: b for b in bookings}
        self.assertEqual(by_name["PRC 940"]["quantity"], 60)
        self.assertEqual(by_name["AM78"]["quantity"], 4)
        self.assertEqual(by_name["Diesel"]["quantity"], 20)
        self.assertEqual(by_name["Generator"]["end"], date(2026, 11, 11))
        # Save/load does not mutate project defaults, and untouched values persist.
        self.assertEqual(load_rules()["items"][0]["reference_quantity"], 20)
        new_items, new_windows = settings_frames(rules)
        self.assertEqual(settings_overrides(defaults, 140, new_items, new_windows), overrides)

    def test_settings_baseline_and_invalid_values(self):
        defaults = load_rules()
        items, windows = settings_frames(defaults)
        rules = load_rules(overrides=settings_overrides(defaults, 100, items, windows))
        self.assertEqual(item_quantity(rules["items"][0], 140), 28)
        items.loc[13, "Baseline quantity"] = -1
        with self.assertRaises(ValueError):
            settings_overrides(defaults, 140, items, windows)

    def test_exact_baseline_and_ceiling(self):
        rules = load_rules()
        for item in rules["items"]:
            q = item["reference_quantity"]
            self.assertEqual(scaled_quantity(item["per_cadet"], 140), q)
            self.assertEqual(scaled_quantity(item["per_cadet"], 280), q * 2)
        self.assertEqual(scaled_quantity("20/140", 141), 21)
        self.assertEqual(scaled_quantity("4/140", 141), 5)
        self.assertEqual(scaled_quantity("70/140", 1), 1)
        self.assertEqual(scaled_quantity("140/140", 141), 141)
        for item in rules["items"]:
            if item["name"] in ("Diesel", "Generator"):
                for strength in (1, 140, 141, 280, 2000):
                    self.assertEqual(item_quantity(item, strength), item["reference_quantity"])
        for invalid in (0, -1, 1.5, True):
            with self.assertRaises(ValueError):
                scaled_quantity("1/140", invalid)

    def test_reference_dates_and_shift(self):
        from datetime import timedelta
        _, bookings, _ = booking_plan(reference_schedule(), 140, load_rules())
        named = {b["name"]: b for b in bookings}
        expected = {
            "KEYHOLE SENSOR": (date(2026, 10, 9), date(2026, 10, 21)),
            "Diesel": (date(2026, 10, 22), date(2026, 11, 10)),
            "Generator": (date(2026, 10, 22), date(2026, 11, 10)),
            "PRC 940": (date(2026, 10, 9), date(2026, 11, 12)),
            "LMG BORELIGHT": (date(2026, 10, 5), date(2026, 10, 19)),
            "NIGHTHAWK": (date(2026, 11, 2), date(2026, 11, 12)),
            "COMPASS MAGNETIC": (date(2026, 9, 28), date(2026, 12, 18)),
        }
        for name, dates in expected.items():
            self.assertEqual((named[name]["start"], named[name]["end"]), dates)
        _, shifted, _ = booking_plan(reference_schedule(35), 140, load_rules())
        for b in shifted:
            self.assertEqual(b["start"], named[b["name"]]["start"] + timedelta(days=35))
            self.assertEqual(b["end"], named[b["name"]]["end"] + timedelta(days=35))

    def test_lancer_never_triggers_equipment(self):
        schedule = pd.DataFrame({"22-Nov-26": ["EX LANCER DAY 7"],
                                 "23-Nov-26": ["EX LANCER DAY 8"]})
        _, bookings, _ = booking_plan(schedule, 140, load_rules())
        self.assertEqual([b["name"] for b in bookings], ["COMPASS MAGNETIC"])

    def test_incomplete_activity_does_not_guess_window(self):
        schedule = pd.DataFrame({"25-Oct-26": ["EX Relentless Warrior"]})
        _, bookings, warnings = booking_plan(schedule, 140, load_rules())
        self.assertNotIn("Diesel", [b["name"] for b in bookings])
        self.assertTrue(any("diesel: incomplete" in w for w in warnings))

    def test_actual_template_multimonth_and_regeneration(self):
        workbook = openpyxl.load_workbook(ROOT / "ocs/template_siao.xlsx")
        original = workbook['(Fill In) SIAO']['C3'].value
        def quantity(row, column):
            cell = ws.cell(row, column)
            for merged in ws.merged_cells.ranges:
                if cell.coordinate in merged:
                    return ws.cell(merged.min_row, merged.min_col).value
            return cell.value
        fill_ce_signals(workbook, reference_schedule(), 140)
        ws = workbook['(Fill In) CE & Signals']
        self.assertEqual(ws.column_dimensions['AN'].width, ws.column_dimensions['AM'].width)
        # September, October, November, December are 36-column blocks.
        self.assertEqual(ws.cell(1, 37).value, date(2026, 10, 1))
        self.assertEqual(ws.cell(61, 36 + 3 + 9).value, 12)
        self.assertIsNone(ws.cell(61, 36 + 3 + 22).value)
        self.assertEqual(ws.cell(116, 36 + 3 + 22).value, 20)
        self.assertEqual(quantity(116, 72 + 3 + 10), 20)
        self.assertIsNone(ws.cell(116, 72 + 3 + 11).value)
        self.assertEqual(ws.cell(127, 36 + 3 + 22).value, 1)
        self.assertEqual(quantity(127, 72 + 3 + 10), 1)
        self.assertIsNone(ws.cell(127, 72 + 3 + 11).value)
        self.assertEqual(quantity(36, 108 + 3 + 18), 140)
        self.assertIsNone(ws.cell(36, 108 + 3 + 19).value)
        self.assertEqual(quantity(36, 72 + 3 + 22), 140)
        self.assertIsNone(ws.cell(13, 72 + 3 + 22).value)
        self.assertIsNone(ws.cell(12, 36 + 3 + 9).value)  # PRC650 was blank.
        self.assertEqual(ws.cell(68, 36 + 3 + 9).value, 7)  # Not template default 1.
        self.assertEqual(ws.cell(76, 36 + 3 + 9).value, 12)  # Not default 3.
        self.assertEqual(workbook['(Fill In) SIAO']['C3'].value, original)
        self.assertEqual(ws.cell(116, 36 + 3 + 22).number_format, '0" L"')
        fill_ce_signals(workbook, reference_schedule(), 141)
        ws = workbook['(Fill In) CE & Signals']
        self.assertEqual(ws.cell(13, 36 + 3 + 9).value, 21)
        self.assertEqual(ws.cell(15, 36 + 3 + 9).value, 5)
        self.assertEqual(ws.cell(127, 36 + 3 + 22).value, 1)
        self.assertEqual(ws.cell(116, 36 + 3 + 22).value, 20)
        self.assertEqual(quantity(36, 108 + 3 + 18), 141)
        self.assertEqual(workbook.sheetnames.count('(Fill In) CE & Signals'), 1)
        self.assertIn('BI127:BS127', {str(r) for r in ws.merged_cells.ranges})
        self.assertIn('BW127:CG127', {str(r) for r in ws.merged_cells.ranges})
        self.assertIsNone(ws.cell(127, 36 + 3 + 23).value)
        workbook.close()


if __name__ == '__main__':
    unittest.main()
