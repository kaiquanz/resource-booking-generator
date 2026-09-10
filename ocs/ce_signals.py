"""Cadet-scaled CE/Signals bookings, shared by both SIAO import routes."""

from calendar import monthrange
from copy import copy
from datetime import datetime, timedelta
from fractions import Fraction
from math import ceil
from pathlib import Path
import re

import yaml
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_RULES_PATH = Path(__file__).with_name("ce_signals.yaml")
HEADER_ROWS = (11, 32, 48, 60, 67, 85, 106, 110, 115, 122)
SECTION_ROWS = (10, 31, 47, 59, 66, 84, 105, 109, 114, 121)
ITEM_ROWS = tuple(
    r for start, end in ((12, 21), (33, 45), (49, 57), (61, 64),
                         (68, 82), (87, 95), (97, 99), (101, 103),
                         (107, 107), (111, 112), (116, 119), (123, 127))
    for r in range(start, end + 1)
)


def load_rules(path=None, overrides=None):
    with open(path or DEFAULT_RULES_PATH, encoding="utf-8") as handle:
        rules = yaml.safe_load(handle)
    if overrides:
        rules["base_cadets"] = overrides.get("base_cadets", rules["base_cadets"])
        for item in rules["items"]:
            changes = overrides.get("items", {}).get(str(item["row"]), {})
            for key in ("reference_quantity", "scaling", "window"):
                if key in changes:
                    item[key] = changes[key]
            item["per_cadet"] = f"{item['reference_quantity']}/{rules['base_cadets']}"
        for name, changes in overrides.get("windows", {}).items():
            if name not in rules["windows"]:
                raise ValueError(f"Unknown CE & Signals window: {name}")
            for key in ("start", "end", "start_offset", "end_offset", "enabled"):
                if key in changes:
                    rules["windows"][name][key] = changes[key]
    return validate_rules(rules)


def validate_rules(rules):
    base = rules["base_cadets"]
    if not isinstance(base, int) or base <= 0:
        raise ValueError("CE & Signals base_cadets must be a positive integer")
    seen = set()
    for item in rules["items"]:
        if not isinstance(item["reference_quantity"], int) or item["reference_quantity"] < 0:
            raise ValueError(f"Quantity must be a non-negative integer: {item['name']}")
        if item.get("scaling", "cadets") not in ("cadets", "fixed"):
            raise ValueError(f"Invalid scaling: {item['name']}")
        if item["row"] in seen or item["row"] not in ITEM_ROWS:
            raise ValueError(f"Invalid/duplicate CE & Signals row: {item['row']}")
        seen.add(item["row"])
        ratio = Fraction(item["per_cadet"])
        if ratio < 0 or ratio * base != item["reference_quantity"]:
            raise ValueError(f"CE & Signals ratio disagrees with baseline: {item['name']}")
        if item["window"] not in rules["windows"]:
            raise ValueError(f"Unknown CE & Signals window: {item['window']}")
    for name, window in rules["windows"].items():
        for key in ("start", "end"):
            if window[key] not in {*rules["anchors"], "course_start", "course_end"}:
                raise ValueError(f"Unknown activity for {name}: {window[key]}")
        for key in ("start_offset", "end_offset"):
            value = window.get(key, 0)
            if not isinstance(value, int) or abs(value) > 365:
                raise ValueError(f"{name}: date offsets must be whole days between -365 and 365")
    return rules


def scaled_quantity(per_cadet, cadet_size):
    """Exact arithmetic avoids floating-point over-rounding at 140 cadets."""
    if isinstance(cadet_size, bool) or int(cadet_size) != cadet_size or cadet_size < 1:
        raise ValueError("Cadet strength must be a positive integer")
    return ceil(Fraction(per_cadet) * int(cadet_size))


def schedule_days(conducts):
    """Read the unfiltered timetable, including unmatched lessons/dekitting."""
    days = {}
    for index, column in enumerate(conducts.columns):
        try:
            day = datetime.strptime(str(column), "%d-%b-%y").date()
        except ValueError:
            continue
        texts = days.setdefault(day, set())
        for value in conducts.iloc[:, index].dropna():
            texts.add(re.sub(r"\s+", " ", str(value)).strip().upper())
    return days


def item_quantity(item, cadet_size):
    if item.get("scaling") == "fixed":
        return item["reference_quantity"]
    return scaled_quantity(item["per_cadet"], cadet_size)


def booking_plan(conducts, cadet_size, rules):
    scaled_quantity("1", cadet_size)
    days = schedule_days(conducts)
    if not days:
        raise ValueError("No dated training events are available for CE & Signals")
    anchors = {"course_start": min(days), "course_end": max(days)}
    for key, spec in rules["anchors"].items():
        matching = sorted(day for day, texts in days.items()
                          if any(re.search(spec["pattern"], text) for text in texts))
        if matching:
            anchors[key] = matching[-1] if spec.get("occurrence") == "last" else matching[0]
    # LANCER is not a booking trigger. Standing compasses remain held throughout.
    lancer_days = {day for day, texts in days.items()
                   if any(re.search(r"^EX\.? LANCER(?: DAY \d+)?$", t) for t in texts)}
    windows, warnings = {}, []
    for name, spec in rules["windows"].items():
        if spec.get("enabled") is False:
            continue
        start = anchors.get(spec["start"])
        end = anchors.get(spec["end"])
        if end is None and spec.get("end_fallback"):
            end = anchors.get(spec["end_fallback"])
            warnings.append("No dekitting found; compasses run to the last timetable date.")
        required = spec.get("requires", [])
        if start is None or end is None or any(key not in anchors for key in required):
            if start is not None or end is not None:
                warnings.append(f"{name}: incomplete activity dates; booking not generated.")
            continue
        start += timedelta(days=spec.get("start_offset", 0))
        end += timedelta(days=spec.get("end_offset", 0))
        if end < start:
            warnings.append(f"{name}: end precedes start; booking not generated.")
            continue
        windows[name] = (start, end)
    bookings = []
    for item in rules["items"]:
        quantity = item_quantity(item, cadet_size)
        if item["window"] not in windows:
            continue
        start, end = windows[item["window"]]
        bookings.append({**item, "quantity": quantity, "start": start, "end": end,
                         "excluded_dates": set() if rules["windows"][item["window"]].get("standing")
                         else lancer_days})
    if rules.get("unassigned_reference_bookings"):
        warnings.append("Late-November radio/MCN quantities are saved but not booked: their activity is unconfirmed and is not LANCER.")
    return days, bookings, warnings


def fill_ce_signals(workbook, conducts, cadet_size, rules_path=None, overrides=None):
    """Rebuild the named monthly output grid, retaining the template's styling.

    Only this output tab is replaced. The legacy photo-layout tab (trailing space)
    is a style source when available; other workbook tabs are left untouched.
    """
    rules = load_rules(rules_path, overrides)
    days, bookings, warnings = booking_plan(conducts, cadet_size, rules)
    title = rules["sheet_name"]
    source = next((workbook[n] for n in (title + " ", title) if n in workbook.sheetnames), None)
    if source is None:
        raise ValueError(f"The SIAO template is missing {title!r}")
    old = workbook[title] if title in workbook.sheetnames else None
    position = workbook.index(old or source)
    ws = workbook.create_sheet("CE Signals generated", position)
    first = min([min(days)] + [b["start"] for b in bookings]).replace(day=1)
    last = max([max(days)] + [b["end"] for b in bookings]).replace(day=1)
    months = []
    month = first
    while month <= last:
        months.append(month)
        month = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
    if len(months) * 36 > 16384:
        raise ValueError("CE & Signals schedule exceeds Excel's column limit")
    for r in range(1, 128):
        ws.row_dimensions[r] = copy(source.row_dimensions[r if r <= 120 else 115])
        ws.row_dimensions[r].index = r
    # The first block supplies all labels/styles; no historical quantities or
    # external/array formulas are carried into the new calendar.
    for block, month in enumerate(months):
        offset = block * 36
        for c in range(1, 37):
            letter = get_column_letter(c + offset)
            dimension = next((d for d in source.column_dimensions.values()
                              if d.min <= c <= d.max), None)
            ws.column_dimensions[letter].width = dimension.width if dimension else 4.78
        for row in range(1, 128):
            for col in range(1, 36):
                src = source.cell(row if row <= 120 else (115 if row == 122 else 116), col)
                dst = ws.cell(row, col + offset)
                dst._style = copy(src._style)
                if col <= 2 and 9 < row <= 120 and isinstance(src.value, (str, int, float)):
                    if not isinstance(src.value, str) or not src.value.startswith("="):
                        dst.value = src.value
        for merged in source.merged_cells.ranges:
            if (merged.max_col <= 35 and merged.max_row <= 120
                    and merged.min_col != 2 and merged.min_row not in ITEM_ROWS):
                ws.merge_cells(start_row=merged.min_row, end_row=merged.max_row,
                               start_column=merged.min_col + offset, end_column=merged.max_col + offset)
        ws.cell(1, offset + 1, month).number_format = "mmm-yy"
        ws.cell(121, offset + 1, "Others")._style = copy(source.cell(114, 1)._style)
        for row, name in {123: "MTW", 124: "Matador Sub Cal (Rounds)",
                          125: "M203 Dummy Rounds", 126: "7.62 Dummy Rounds", 127: "Generator"}.items():
            ws.cell(row, offset + 1, name)
        for row in HEADER_ROWS:
            ws.cell(row, offset + 2, "Default")
        for row in ITEM_ROWS:
            ws.cell(row, offset + 2).value = None
        count = monthrange(month.year, month.month)[1]
        for i in range(33):
            day = month + timedelta(days=i - 1)
            valid = i <= count + 1
            for row in HEADER_ROWS:
                cell = ws.cell(row, offset + 3 + i)
                cell.value = day if valid else None
                cell.number_format = "dd"
                cell.fill = PatternFill("solid", fgColor="FFF200")
            for row in ITEM_ROWS:
                cell = ws.cell(row, offset + 3 + i)
                cell.fill = PatternFill("solid", fgColor="E6C7C7" if valid and day.weekday() >= 5 else "FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for item in rules["items"]:
            row = item["row"]
            ws.cell(row, offset + 1, item["name"])
            base_cell = ws.cell(row, offset + 2, item_quantity(item, cadet_size))
            base_cell.number_format = '0" L"' if item.get("unit") == "L" else "0"
            calculation = ("Fixed quantity, independent of cadet strength. "
                           if item.get("scaling") == "fixed" else
                           f"Round up after multiplying by {cadet_size}. ")
            base_cell.comment = Comment(
                f"Reference: {item['reference_quantity']} / {rules['base_cadets']} cadets. "
                f"Stored ratio: {item['per_cadet']}. {calculation}"
                f"Booking rule: {item['window']}. See ce_signals.yaml for date assumptions.", "TP Automation")
        for booking in bookings:
            run_start = None
            for i in range(count + 2):
                day = month + timedelta(days=i - 1)
                if booking["start"] <= day <= booking["end"] and day not in booking["excluded_dates"]:
                    cell = ws.cell(booking["row"], offset + 3 + i, booking["quantity"])
                    cell.number_format = '0" L"' if booking.get("unit") == "L" else "0"
                    if run_start is None:
                        run_start = cell.column
                elif run_start is not None:
                    if offset + 2 + i > run_start:
                        ws.merge_cells(start_row=booking["row"], end_row=booking["row"],
                                       start_column=run_start, end_column=offset + 2 + i)
                    run_start = None
            if run_start is not None and offset + count + 4 > run_start:
                ws.merge_cells(start_row=booking["row"], end_row=booking["row"],
                               start_column=run_start, end_column=offset + count + 4)
    ws.freeze_panes = "C12"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 0
    ws.page_setup.fitToHeight = 1
    ws.print_area = f"A1:{get_column_letter(len(months) * 36 - 1)}127"
    if old is not None:
        workbook.remove(old)
    ws.title = title
    return [{k: v for k, v in b.items() if k != "excluded_dates"} for b in bookings], warnings
